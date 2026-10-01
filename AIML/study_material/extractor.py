from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from parser.cleaner import TextCleaner


class StudyMaterialExtractor:

    @staticmethod
    def extract_text(file_path):
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        # TXT
        if path.suffix.lower() == ".txt":
            with open(path, "r", encoding="utf-8") as file:
                return file.read()

        # PDF
        elif path.suffix.lower() == ".pdf":
            reader = PdfReader(str(path))
            text = ""

            for page in reader.pages:
                page_text = page.extract_text()

                if page_text:
                    text += page_text + "\n"

            return text

        # DOCX
        elif path.suffix.lower() == ".docx":
            document = Document(str(path))
            text = ""

            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"

            return text

        # PPTX
        elif path.suffix.lower() == ".pptx":
            presentation = Presentation(str(path))
            text = ""

            for slide in presentation.slides:
                for shape in slide.shapes:

                    if hasattr(shape, "text") and shape.text.strip():
                        text += shape.text + "\n"

            return TextCleaner.clean(text)

        else:
            raise ValueError(
                "Unsupported file type. "
                "Please use .txt, .pdf, .docx, or .pptx"
            )


if __name__ == "__main__":

    file_path = input("Enter study material path: ")

    text = StudyMaterialExtractor.extract_text(file_path)

    print("\n========== EXTRACTED TEXT ==========")
    print(text)