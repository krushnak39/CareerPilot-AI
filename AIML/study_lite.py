"""
Lightweight study-material assistant (no LLM, no GPU, no downloads).

Works on pdf / docx / pptx / txt and provides:
    ask        - most relevant passages for a question (BM25-style ranking)
    summary    - the most informative sentences
    flashcards - definition-based term / meaning cards
    quiz       - multiple-choice questions built from definitions
"""

import math
import random
import re
from collections import Counter
from pathlib import Path

STOPWORDS = set(
    """a an and are as at be by for from has have in is it its of on or that the
    to was were will with this these those which who what when where how why do
    does did can could should would about into than then there their they them
    you your we our not but if so such also may might""".split()
)


# ------------------------------------------------------------------
# Text extraction
# ------------------------------------------------------------------
def extract_text(path):
    suffix = Path(path).suffix.lower()

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        return "\n".join((p.extract_text() or "") for p in reader.pages)

    if suffix == ".docx":
        from docx import Document

        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)

    if suffix == ".pptx":
        from pptx import Presentation

        parts = []
        for slide in Presentation(path).slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    parts.append(shape.text_frame.text)
        return "\n".join(parts)

    if suffix in (".txt", ".md"):
        return Path(path).read_text(encoding="utf-8", errors="ignore")

    raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")


# ------------------------------------------------------------------
# Helpers
# ------------------------------------------------------------------
def _clean(text):
    return re.sub(r"[ \t]+", " ", text.replace("\r", "")).strip()


def split_sentences(text):
    text = re.sub(r"\n{2,}", ". ", text)
    text = text.replace("\n", " ")
    parts = re.split(r"(?<=[.!?])\s+", text)
    return [p.strip() for p in parts if len(p.strip()) >= 25]


def tokens(text):
    return [
        t
        for t in re.findall(r"[a-zA-Z0-9+#.]+", text.lower())
        if t not in STOPWORDS and len(t) > 1
    ]


def make_chunks(sentences, size=3):
    return [
        " ".join(sentences[i : i + size]) for i in range(0, len(sentences), size)
    ]


# ------------------------------------------------------------------
# Features
# ------------------------------------------------------------------
def ask(text, question, k=3):
    sentences = split_sentences(_clean(text))
    chunks = make_chunks(sentences, 3)
    if not chunks:
        return {"answer": "No readable text found in this material.", "passages": []}

    q_terms = tokens(question)
    if not q_terms:
        return {"answer": "Please ask a more specific question.", "passages": []}

    docs = [Counter(tokens(c)) for c in chunks]
    n = len(docs)
    avg_len = sum(sum(d.values()) for d in docs) / n or 1
    df = Counter(t for d in docs for t in d)

    scored = []
    for chunk, d in zip(chunks, docs):
        length = sum(d.values()) or 1
        score = 0.0
        for t in q_terms:
            if t in d:
                idf = math.log(1 + (n - df[t] + 0.5) / (df[t] + 0.5))
                tf = d[t]
                score += idf * (tf * 2.2) / (tf + 1.2 * (0.25 + 0.75 * length / avg_len))
        scored.append((score, chunk))

    scored.sort(key=lambda x: x[0], reverse=True)
    top = [c for s, c in scored[:k] if s > 0]

    if not top:
        return {
            "answer": "I could not find this information in the provided study material.",
            "passages": [],
        }

    return {"answer": top[0], "passages": top}


def summary(text, n=5):
    sentences = split_sentences(_clean(text))
    if not sentences:
        return {"summary": []}

    freq = Counter(t for s in sentences for t in tokens(s))
    scored = []
    for i, s in enumerate(sentences):
        words = tokens(s)
        if not words:
            continue
        score = sum(freq[w] for w in words) / (len(words) ** 0.8)
        scored.append((score, i, s))

    best = sorted(scored, key=lambda x: x[0], reverse=True)[:n]
    best.sort(key=lambda x: x[1])  # keep original order
    return {"summary": [s for _, _, s in best]}


DEFINITION = re.compile(
    r"^(?P<term>[A-Z][\w\s\-/()+.#]{1,50}?)\s+"
    r"(?:is|are|refers to|means|is defined as|is used to)\s+(?P<meaning>.{12,260}?)[.!?]?$"
)


def _definitions(text):
    found = []
    seen = set()
    for sentence in split_sentences(_clean(text)):
        m = DEFINITION.match(sentence)
        if not m:
            continue
        term = m.group("term").strip()
        if len(term.split()) > 6 or term.lower() in seen:
            continue
        seen.add(term.lower())
        found.append((term, sentence))
    return found


def flashcards(text, n=10):
    cards = [
        {"front": f"What is {term}?", "back": sentence}
        for term, sentence in _definitions(text)[:n]
    ]
    return {"flashcards": cards}


def quiz(text, n=5, seed=None):
    defs = _definitions(text)
    rng = random.Random(seed)
    questions = []

    terms = [t for t, _ in defs]
    for term, sentence in defs[:n]:
        blanked = re.sub(re.escape(term), "_____", sentence, count=1, flags=re.I)
        distractors = [t for t in terms if t != term]
        rng.shuffle(distractors)
        options = [term] + distractors[:3]
        rng.shuffle(options)
        questions.append(
            {
                "question": blanked,
                "options": options,
                "answer": term,
            }
        )
    return {"quiz": questions}


def run(path, action, question="", n=None):
    text = extract_text(path)
    if not text.strip():
        raise ValueError("No readable text found in this file.")

    if action == "ask":
        return ask(text, question)
    if action == "summary":
        return summary(text, n or 5)
    if action == "flashcards":
        return flashcards(text, n or 10)
    if action == "quiz":
        return quiz(text, n or 5)
    raise ValueError("Unknown action")
